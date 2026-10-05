from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Pt

base = Path(__file__).resolve().parents[1]
old = Document(base / 'final_report.docx')
doc = Document(base.parent / 'template_reference.docx')

def fill(p, text):
    props = deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    for child in list(p._p):
        if child.tag.endswith('}r'):
            p._p.remove(child)
    r = p.add_run(text)
    if props is not None:
        r._r.insert(0, props)

mapping = {0:0, 1:1, 7:7, 12:12, 16:16, 21:23, 25:27, 31:33, 32:34}
for dst, src in mapping.items():
    fill(doc.paragraphs[dst], old.paragraphs[src].text)
fill(doc.paragraphs[0], 'Tabular Classification')
fill(doc.paragraphs[8], '')
fill(doc.paragraphs[17], '')
fill(doc.paragraphs[33], '\n'.join(p.text for p in old.paragraphs if p.style.name == 'VITA'))

# Insert the comparison table after the Results paragraph, in the body flow.
doc.paragraphs[16]._p.addnext(deepcopy(old.tables[0]._tbl))

# The template footer is a three-cell table. Replace only its text runs.
seen = set()
for section in doc.sections:
    for footer in (section.footer, section.first_page_footer, section.even_page_footer):
        if footer.part.partname in seen:
            continue
        seen.add(footer.part.partname)
        for table in footer.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        for run in p.runs:
                            run.text = run.text.replace('[Document title]', 'Tabular Classification').replace('[Author name]', 'LIN TONG')

doc.save(base / 'final_report.docx')
print('Rebuilt from the original template with its footer table intact')
