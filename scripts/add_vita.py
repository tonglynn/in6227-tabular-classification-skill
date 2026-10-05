from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from copy import deepcopy
from docx.shared import Pt

base = Path(__file__).resolve().parents[1]
path = base / 'final_report.docx'
doc = Document(path)
for p in list(doc.paragraphs):
    if p.text.startswith(('Name:', 'Matric number:', 'Assignment:', 'Model (LLM):', 'LLM interface:', 'GitHub:')):
        p._element.getparent().remove(p._element)
anchor = next(p for p in doc.paragraphs if p.style.name == 'Body Text')
values = [
    'Name: LIN TONG', 'Matric number: G2608825L',
    'Assignment: IN6227-2023-Assignment-1, Variant 2',
    'Model (LLM): GLM-5.2', 'LLM interface: TraeCode (Codex)',
    'GitHub: https://github.com/tonglynn/in6227-tabular-classification-skill']
for value in values:
    p = anchor.insert_paragraph_before(value, style='VITA')
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in p.runs:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(10)
# The original VITA paragraph terminates the two-column section. Preserve
# that boundary after the final metadata line, not before the metadata.
reference = Document(base.parent / 'template_reference.docx')
section = reference.paragraphs[33]._p.pPr.sectPr
p._p.get_or_add_pPr().append(deepcopy(section))
doc.save(path)
print('vita restored')
