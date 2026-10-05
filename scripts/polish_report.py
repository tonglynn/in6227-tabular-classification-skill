from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement

base = Path(__file__).resolve().parents[1]
path = base / 'final_report.docx'
doc = Document(path)

# A Word table is allowed to split at a column break. Keep the comparison in
# a compact Results sentence so the original two-column flow stays intact.
table = doc.tables[0]
table._element.getparent().remove(table._element)
results = next(p for p in doc.paragraphs if p.text.startswith('Test-set performance'))
results.add_run(' CV comparison: dummy F1=0.235, logistic regression F1=0.660, '
                'random forest F1=0.664, gradient boosting F1=0.626.')

# Restore VITA as separate template paragraphs instead of one tabbed block.
vita = next(p for p in doc.paragraphs if p.style.name == 'VITA')
values = [
    'Name: LIN TONG',
    'Matric number: G2608825L',
    'Assignment: IN6227-2023-Assignment-1, Variant 2',
    'Model (LLM): GLM-5.2',
    'LLM interface: TraeCode (Codex)',
    'GitHub: https://github.com/tonglynn/in6227-tabular-classification-skill',
]
vita.text = values[0]
vita.alignment = WD_ALIGN_PARAGRAPH.LEFT
# Remove any prior generated VITA lines before rebuilding them.
for p in list(doc.paragraphs):
    if p._element is not vita._element and p.style.name == 'VITA':
        p._element.getparent().remove(p._element)
anchor = vita._p
for value in values[1:]:
    newp = OxmlElement('w:p')
    newp.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rsidR', '00C60E6A')
    anchor.addnext(newp)
    anchor = newp
    # Use a normal VITA paragraph with the same style ID as the original.
    ppr = OxmlElement('w:pPr')
    style = OxmlElement('w:pStyle')
    style.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val', 'VITA')
    ppr.append(style)
    jc = OxmlElement('w:jc')
    jc.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val', 'left')
    ppr.append(jc)
    newp.append(ppr)
    r = OxmlElement('w:r')
    t = OxmlElement('w:t')
    t.text = value
    r.append(t)
    newp.append(r)

doc.save(path)
print('polished')
